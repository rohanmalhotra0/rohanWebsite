"""Build every résumé in content.py as a one-page PDF and an editable DOCX.

    python3 resumes/src/build.py

PDFs are rendered with the pre-installed Chromium (via print.mjs) and auto-fit to
exactly one page. DOCX files are checked with LibreOffice and shrunk until they
also fit on one page.
"""

import html
import json
import math
import pathlib
import re
import shutil
import subprocess
import sys
import tempfile

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

from content import CONTACT, NAME, RESUMES

SRC = pathlib.Path(__file__).resolve().parent
ROOT = SRC.parent
PDF_DIR = ROOT / "pdf"
DOCX_DIR = ROOT / "docx"
BUILD = ROOT / "build"

MARGIN_X, MARGIN_Y = 0.55, 0.45  # inches
FILE_PREFIX = NAME.replace(" ", "_") + "_Resume_"


def segments(text):
    """Split '**bold** plain' markup into (text, style) pairs."""
    parts = re.split(r"(\*\*.+?\*\*)", text)
    return [(p[2:-2], "b") if p.startswith("**") else (p, "") for p in parts if p]


# ---------------------------------------------------------------- HTML

CSS = (SRC / "resume.css").read_text()


def h_segs(segs):
    out = []
    for seg in segs:
        text, style, link = (seg + (None,))[:3]
        t = html.escape(text)
        if style == "b":
            t = f"<b>{t}</b>"
        elif style == "i":
            t = f"<i>{t}</i>"
        if link:
            t = f'<a href="{html.escape(link)}">{t}</a>'
        out.append(t)
    return "".join(out)


def h_rich(text):
    return h_segs(segments(text))


def render_html(blocks):
    contact = '<span class="sep">|</span>'.join(
        h_segs([(t, "", l)] if l else [(t, "")]) for t, l in CONTACT
    )
    out = [f'<header><h1 class="name">{html.escape(NAME)}</h1><div class="contact">{contact}</div></header>']
    for block in blocks:
        if "para" in block:
            out.append(f'<p class="summary">{h_rich(block["para"])}</p>')
            continue
        out.append(f'<section><h2>{html.escape(block["title"])}</h2>')
        for b in block["blocks"]:
            if "lines" in b:
                out.append('<div class="entry lines">' + "".join(
                    f"<p><b>{html.escape(k)}:</b> {h_rich(v)}</p>" for k, v in b["lines"]) + "</div>")
                continue
            out.append('<div class="entry">')
            for g in b["groups"]:
                out.append('<div class="group">')
                for row in g["rows"]:
                    out.append(f'<div class="row"><span class="l">{h_segs(row["left"])}</span>'
                               f'<span class="r">{h_segs(row["right"])}</span></div>')
                if g["bullets"]:
                    out.append("<ul>" + "".join(f"<li>{h_rich(x)}</li>" for x in g["bullets"]) + "</ul>")
                out.append("</div>")
            out.append("</div>")
        out.append("</section>")
    body = "\n".join(out)
    return (f'<!doctype html><html lang="en"><head><meta charset="utf-8"><title>{html.escape(NAME)} Résumé</title>'
            f"<style>{CSS}</style></head><body><div class=\"page\"><div class=\"content\">{body}</div></div></body></html>")


# ---------------------------------------------------------------- DOCX

FONT = "Cambria"  # LibreOffice measures it with metric-identical Caladea


def set_run_font(run, size, style=""):
    run.font.name = FONT
    run._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)
    run.font.size = Pt(size)
    run.bold = "b" in style
    run.italic = "i" in style


def add_hyperlink(par, text, url, size, style=""):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink",
                          is_external=True)
    link = OxmlElement("w:hyperlink")
    link.set(qn("r:id"), r_id)
    run = par.add_run(text)
    set_run_font(run, size, style)
    link.append(run._element)
    par._p.append(link)


def add_segs(par, segs, size):
    for seg in segs:
        text, style, link = (seg + (None,))[:3]
        if link:
            add_hyperlink(par, text, link, size, style)
        else:
            set_run_font(par.add_run(text), size, style)


def bottom_border(par):
    pbdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    for k, v in (("w:val", "single"), ("w:sz", "6"), ("w:space", "1"), ("w:color", "111111")):
        bottom.set(qn(k), v)
    pbdr.append(bottom)
    par._p.get_or_add_pPr().append(pbdr)


def fmt(par, before=0.0, after=0.0, line=1.0):
    pf = par.paragraph_format
    pf.space_before = Pt(before)
    pf.space_after = Pt(after)
    pf.line_spacing = line
    pf.widow_control = False


def render_docx(blocks, fs, gap, path):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = Inches(MARGIN_X)
    sec.top_margin = sec.bottom_margin = Inches(MARGIN_Y)
    width = Inches(8.5 - 2 * MARGIN_X)
    normal = doc.styles["Normal"]
    normal.font.name = FONT
    normal.font.size = Pt(fs)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), FONT)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt(p)
    set_run_font(p.add_run(NAME), fs * 2.1, "b")

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fmt(p, before=1.5 * gap)
    for i, (t, link) in enumerate(CONTACT):
        if i:
            set_run_font(p.add_run("  |  "), fs * 0.95)
        add_segs(p, [(t, "", link)] if link else [(t, "")], fs * 0.95)

    def row_par(row, before):
        par = doc.add_paragraph()
        fmt(par, before=before)
        par.paragraph_format.tab_stops.add_tab_stop(width, WD_TAB_ALIGNMENT.RIGHT)
        add_segs(par, row["left"], fs)
        set_run_font(par.add_run("\t"), fs)
        add_segs(par, row["right"], fs)

    def bullet_par(text):
        par = doc.add_paragraph()
        fmt(par, before=0.5 * gap)
        pf = par.paragraph_format
        pf.left_indent = Inches(0.2)
        pf.first_line_indent = Inches(-0.13)
        pf.tab_stops.add_tab_stop(Inches(0.2))
        set_run_font(par.add_run("•\t"), fs)
        add_segs(par, segments(text), fs)

    for block in blocks:
        if "para" in block:
            par = doc.add_paragraph()
            fmt(par, before=5 * gap)
            add_segs(par, segments(block["para"]), fs)
            continue
        h = doc.add_paragraph()
        fmt(h, before=7 * gap, after=2.5 * gap)
        bottom_border(h)
        r = h.add_run(block["title"])
        set_run_font(r, fs * 1.05, "b")
        r.font.all_caps = True
        for bi, b in enumerate(block["blocks"]):
            entry_before = 0 if bi == 0 else 3.5 * gap
            if "lines" in b:
                for li, (k, v) in enumerate(b["lines"]):
                    par = doc.add_paragraph()
                    fmt(par, before=(0.8 * gap) if li else entry_before)
                    add_segs(par, [(k + ": ", "b")] + segments(v), fs)
                continue
            for gi, g in enumerate(b["groups"]):
                for ri, row in enumerate(g["rows"]):
                    first = gi == 0 and ri == 0
                    row_par(row, entry_before if first else (2 * gap if ri == 0 else 0))
                for x in g["bullets"]:
                    bullet_par(x)
    doc.save(path)


def docx_pages(paths):
    """Convert DOCX files to PDF with LibreOffice and return page counts."""
    with tempfile.TemporaryDirectory() as tmp:
        subprocess.run(["soffice", "--headless", "--convert-to", "pdf", "--outdir", tmp, *map(str, paths)],
                       check=True, capture_output=True)
        counts = {}
        for p in paths:
            info = subprocess.run(["pdfinfo", str(pathlib.Path(tmp) / (p.stem + ".pdf"))],
                                  check=True, capture_output=True, text=True).stdout
            counts[p] = int(re.search(r"Pages:\s+(\d+)", info).group(1))
        return counts


# ---------------------------------------------------------------- main

def main():
    only = set(sys.argv[1:])
    names = [n for n in RESUMES if not only or n in only]
    BUILD.mkdir(exist_ok=True)
    PDF_DIR.mkdir(exist_ok=True)
    DOCX_DIR.mkdir(exist_ok=True)

    jobs = []
    for name in names:
        html_path = BUILD / f"{name}.html"
        html_path.write_text(render_html(RESUMES[name]))
        jobs.append({"name": name, "html": str(html_path), "pdf": str(PDF_DIR / f"{FILE_PREFIX}{name}.pdf")})
    jobs_path = BUILD / "jobs.json"
    jobs_path.write_text(json.dumps(jobs))
    subprocess.run(["node", str(SRC / "print.mjs"), str(jobs_path)], check=True)
    report = json.loads((BUILD / "report.json").read_text())

    # DOCX: Word font sizes snap to half points, so start at the PDF size rounded
    # down and shrink spacing, then font, until LibreOffice says one page.
    settings = {r["name"]: [math.floor(r["fs"] * 2) / 2, r["gap"]] for r in report}
    pending = list(names)
    for _ in range(12):
        paths = {}
        for name in pending:
            fs, gap = settings[name]
            paths[name] = DOCX_DIR / f"{FILE_PREFIX}{name}.docx"
            render_docx(RESUMES[name], fs, gap, paths[name])
        counts = docx_pages(list(paths.values()))
        pending = [n for n in pending if counts[paths[n]] > 1]
        if not pending:
            break
        for name in pending:
            fs, gap = settings[name]
            settings[name] = [fs, gap - 0.15] if gap > 0.7 else [fs - 0.5, 1.0]
    for r in report:
        r["docx_fs"], r["docx_gap"] = settings[r["name"]]
    if pending:
        print("DOCX still over one page:", pending)

    print(f"\n{'résumé':24} {'pdf pt':>7} {'gap':>5} {'docx pt':>8}  notes")
    for r in report:
        notes = "; ".join(r["warnings"]) or "ok"
        print(f"{r['name']:24} {r['fs']:7.2f} {r['gap']:5.2f} {r['docx_fs']:8.1f}  {notes}")
    if any(r["pages"] != 1 for r in report):
        sys.exit("A PDF is not exactly one page.")


if __name__ == "__main__":
    main()
